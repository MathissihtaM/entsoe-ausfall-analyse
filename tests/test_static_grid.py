from pathlib import Path

from src.static_grid import (
    aggregate_by_eic,
    compact_eic,
    discover_zip_url,
    load_grid_elements,
    lookup_grid,
    make_fixture_workbook,
    nominal_mw,
    parse_workbook,
    write_grid_csv,
)


def test_nominal_mw_sqrt3_and_snom():
    assert abs(nominal_mw(380, 2500) - 1645.448) < 0.01
    assert nominal_mw(380, 1000, snom=400) == 400
    assert nominal_mw(None, 1000) is None
    assert nominal_mw(380, 0) is None


def test_parse_and_aggregate_double_circuit(tmp_path: Path):
    xlsx = tmp_path / "sgm.xlsx"
    xlsx.write_bytes(make_fixture_workbook([
        {
            "sheet": "Lines",
            "name": "Röhrsdorf - Hradec 551",
            "eic": "11TD8L551------X",
            "tso": "50HERTZ",
            "sub_from": "Röhrsdorf",
            "sub_to": "Hradec",
            "kv": 380,
            "imax": 2000,
        },
        {
            "sheet": "Lines",
            "name": "Röhrsdorf - Hradec 552",
            "eic": "11TD8L551------X",
            "tso": "50HERTZ",
            "sub_from": "Röhrsdorf",
            "sub_to": "Hradec",
            "kv": 380,
            "imax": 2000,
        },
        {
            "sheet": "Tielines",
            "name": "Kontek",
            "eic": "11T0-0000-0265-C",
            "tso": "50Hertz",
            "sub_from": "Bentwisch",
            "sub_to": "Bjæverskov",
            "kv": 400,
            "imax": 600,
        },
        {
            "sheet": "Lines",
            "name": "Foreign only",
            "eic": "10T-CZ-SK-00005Q",
            "tso": "CEPS",
            "sub_from": "Hradec",
            "sub_to": "Sokolnice",
            "kv": 380,
            "imax": 2000,
        },
        {
            "sheet": "Transformers",
            "name": "UW Röhrsdorf T1",
            "eic": "11TD8TROH------1",
            "tso": "50Hertz",
            "sub_from": "Röhrsdorf",
            "sub_to": "",
            "kv": 380,
            "imax": 1500,
            "snom": 300,
        },
    ]))
    rows = parse_workbook(xlsx)
    eics = {r["key"] for r in rows}
    assert compact_eic("11TD8L551------X") in eics
    assert compact_eic("11T0-0000-0265-C") in eics
    assert compact_eic("10T-CZ-SK-00005Q") not in eics
    trafo = next(r for r in rows if r["type"] == "Transformer")
    assert trafo["sub_to"] == "Röhrsdorf"
    assert trafo["nominal_mw"] == 300

    agg = aggregate_by_eic(rows)
    line = next(r for r in agg if r["key"] == compact_eic("11TD8L551------X"))
    assert line["imax"] == 4000
    assert abs(line["nominal_mw"] - nominal_mw(380, 4000)) < 0.01


def test_lookup_and_csv_roundtrip(tmp_path: Path):
    dest = tmp_path / "grid.csv"
    write_grid_csv([
        {
            "eic": "11TD8L551------X",
            "name": "Röhrsdorf - Hradec",
            "tso": "50Hertz",
            "type": "Line",
            "sub_from": "Röhrsdorf",
            "sub_to": "Hradec",
            "kv": 380,
            "nominal_mw": 1645.4,
            "key": compact_eic("11TD8L551------X"),
        }
    ], dest)
    rows = load_grid_elements(dest)
    table = {r["key"]: r for r in rows}
    hit = lookup_grid("11TD8L551------X", table)
    assert hit is not None
    assert hit["sub_from"] == "Röhrsdorf"
    assert lookup_grid("11W0-0000-0026-Y", table) is None


def test_parse_real_two_row_headers(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Lines"
    ws.append([None, None, None, "Substation_1", "Substation_2", None, "Maximum Current Imax (A)", None, None, None, None, None, None])
    ws.append(["NE_name", "EIC_Code", "TSO", "Full_name", "Full_name", "Voltage_level(kV)", "Period 1", "Period 2", "Period 3", "Period 4", "Period 5", "Period 6", "Fixed"])
    ws.append(["Röhrsdorf - Hradec 445", "10T-CZ-DE-00002U", "50HERTZ", "Röhrsdorf", "Hradec Vychod", 380, None, None, None, None, None, None, 2520])
    ws.append(["SETZE N", "11T0-0000-0117-Z", "AMPRION GMBH", "Altenkleusheim", "Setzer Wiese", 220, "", "", "", "", "", "", "", None, 740])
    # extra DLR columns after Fixed
    ws.cell(1, 14, "Dynamic line rating (DLR)")
    ws.cell(2, 14, "DLRmin(A)")
    ws.cell(2, 15, "DLRmax(A)")
    dest = tmp_path / "real.xlsx"
    wb.save(dest)
    rows = parse_workbook(dest)
    assert len(rows) == 2
    first = next(r for r in rows if r["eic"].startswith("10T"))
    dlr = next(r for r in rows if r["eic"].startswith("11T0"))
    assert first["sub_from"] == "Röhrsdorf"
    assert first["imax"] == 2520
    assert first["kv"] == 380
    assert first["nominal_mw"] == nominal_mw(380, 2520)
    assert dlr["imax"] == 740
    assert dlr["nominal_mw"] == nominal_mw(220, 740)


def test_discover_zip_skips_outdated():
    html = """
    <a href="/files/outdated_202503_Core Static Grid Model_7th release.zip">old</a>
    <a href="/files/202509_Core Static Grid Model_8th release.zip">eight</a>
    <a href="/files/202603_Core Static Grid Model_9th release.zip">nine</a>
    """
    url = discover_zip_url(html, "https://www.jao.eu/static-grid-model")
    assert "202603" in url
    assert "outdated" not in url.lower()
